"""보고서 생성 모듈.

1) 주간 자동보고서(weekly_report_bytes) — 조원이 만든 NCR 양식(utils/ncr_report.py)을 그대로
   재사용한다. 케이스 전용 섹션(즉시조치/원인분석/시정조치계획/효과성검증/승인)은 주간 집계에는
   해당 사항이 없어 빈칸으로 나가지만, 문서 구조·표 개수·섹션 순서는 NCR과 완전히 동일하다.
   자동보고서 탭(tab4) 화면의 미리보기도 weekly_report_summary() 하나에서 값을 받아 그린다 —
   화면과 다운로드 파일이 따로 계산하면 내용이 어긋나므로 반드시 이 함수 하나만 쓸 것.

2) 오늘의 현황 보고서(today_status_report_bytes) — 오늘의 현황 탭(tab3_spc.py)의
   '🔎 위험도순 확인' 패널에서 받는 보고서. 그 화면에 실제로 떠 있는 값만 담는다
   (처리현황 · 핵심 지표 · 검사자 처리 결과 · 위험도순 대기 목록 · SPC 관리도 상태).
   값은 전부 tab3에서 인자로 넘겨받는다 — 여기서 tab3를 import하면 순환 import가 된다.
"""
import io
from datetime import date, datetime, timedelta

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from utils.ncr_report import _apply_korean, _bold_line, _cell, _shade_cell, build_ncr_docx

NAVY_HEX = "184F95"          # NCR 양식과 같은 남색
STRIPE_HEX = "F2F4F7"        # 짝수 행 옅은 회색 줄무늬
OK_RGB = RGBColor(0x0C, 0xA3, 0x0C)
BAD_RGB = RGBColor(0xD0, 0x3B, 0x3B)

STATUS_LABELS = {
    "auto_pass": "자동 통과",
    "auto_reject": "자동 배출",
    "attention": "사람 확인 필요",
    "attention_crack": "사람 확인 필요 · 균열계열 의심",
    "attention_margin": "사람 확인 필요 · 균열/용입불량 경계 모호",
}


# ======================================================================
# 1) 주간 자동보고서 (NCR 양식)
# ======================================================================
def weekly_report_summary(thresholds):
    """주간 보고서에 들어가는 값을 한 곳에서 계산한다. tab4 미리보기와 docx가 같이 쓴다."""
    from utils.dummy_data import load_validation_predictions
    from utils.routing import compute_kpis

    df = load_validation_predictions()
    kpis = compute_kpis(df, thresholds)

    period_end = date.today()
    period_start = period_end - timedelta(days=6)
    total_count = 3214  # TODO: 실제 이번 주 검사 물량 집계로 교체
    human_review_count = round(total_count * kpis["human_review_rate"])
    margin_hold_count = round(total_count * kpis["margin_hold_rate"])
    d1_miss_count = round(total_count * kpis["d1_miss_rate"])
    d4_miss_count = round(total_count * kpis["d4_miss_rate"])

    # NCR 양식은 케이스 1건 기준 칸이라, 주간 집계 값을 최대한 그 칸 의미에 맞춰 채운다.
    # - 제품/로트/필름 번호 칸 → 이번 주 총 검사 물량
    # - 부적합 유형 칸 → 균열(D1)/용입불량(D4) 중 미검출 건수가 더 많은 쪽 (참고용 표시)
    # - AI 판정 확률 칸 → 이번 주 자동화율
    # - 검사원 최종 판정 칸 → "주간 집계 보고"
    dominant = "균열(D1)" if d1_miss_count >= d4_miss_count else "용입불량(D4)"

    return {
        "period_start": period_start,
        "period_end": period_end,
        "total_count": total_count,
        "human_review_count": human_review_count,
        "margin_hold_count": margin_hold_count,
        "d1_miss_count": d1_miss_count,
        "d4_miss_count": d4_miss_count,
        "automation_rate": kpis["automation_rate"],
        "dominant": dominant,
        "decision": "주간 집계 보고",
        "report_no": f"WEEKLY-{period_end.strftime('%Y%m%d')}",
        "product_text": (
            f"이번 주 전체 검사 물량 {total_count:,}건 "
            f"(사람확인 {human_review_count:,}건 · Margin보류 {margin_hold_count:,}건)"
        ),
    }


def weekly_report_bytes(thresholds):
    s = weekly_report_summary(thresholds)
    docx_bytes = build_ncr_docx(
        image_id=s["product_text"],
        dominant_class=s["dominant"],
        ai_prob=s["automation_rate"],
        decision=s["decision"],
        report_no=s["report_no"],
        found_date=s["period_end"],
    )
    return docx_bytes, f"weekly_report_{s['period_end'].isoformat()}.docx"


# ======================================================================
# 2) 오늘의 현황 보고서 (tab3 위험도순 확인 패널 전용)
# ======================================================================
def _head_row(table, headers):
    for c, head in enumerate(headers):
        cell = table.cell(0, c)
        _shade_cell(cell, NAVY_HEX)
        _cell(cell, head, bold=True, align_center=True)
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def _body_cell(table, r, c, text, bold=False, color=None, center=True):
    cell = table.cell(r, c)
    _cell(cell, text, bold=bold, align_center=center)
    if color is not None:
        cell.paragraphs[0].runs[0].font.color.rgb = color
    if r % 2 == 0:  # 헤더(0행) 다음부터 짝수 행 줄무늬
        _shade_cell(cell, STRIPE_HEX)


def _table(doc, rows, cols, widths=None):
    """widths: 열 너비(pt) 목록. 본문 폭(여백 제외 약 528pt)에 맞춰 지정한다.
    Word는 셀마다 너비를 따로 가지므로 모든 행의 셀에 같은 값을 넣어야 적용된다."""
    t = doc.add_table(rows=rows, cols=cols)
    t.style = "Table Grid"
    if widths:
        t.autofit = False
        for c, w in enumerate(widths):  # 표 격자(gridCol) 너비 — LibreOffice 등은 이 값을 따른다
            t.columns[c].width = Pt(w)
        for row in t.rows:
            for c, w in enumerate(widths):
                row.cells[c].width = Pt(w)
    return t


def _plain_line(doc, text, size=9, italic=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    _apply_korean(run.font, run._element, size=Pt(size))
    run.italic = italic
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    return p


def today_status_report_bytes(routing_summary, core_kpis, open_items, approved_n, rejected_n,
                              ncr_list, spc, thresholds):
    """
    routing_summary : tab3 ROUTING_SUMMARY (label/n/pct/status)
    core_kpis       : tab3 CORE_KPIS [(이름, 값, 설명), ...] — 0 자동화율, 1 D1 미검출, 2 D4 미검출
    open_items      : 위험도순 확인 대기 목록 (화면에 떠 있는 순서 그대로)
    ncr_list        : [(image_id, 파일명), ...] 오늘 발행된 NCR
    spc             : {"date","value","cl","ucl","breached","n_outliers","demo"}
    """
    now = datetime.now()
    report_no = f"DAILY-{now.strftime('%Y%m%d')}"

    doc = Document()
    normal = doc.styles["Normal"]
    _apply_korean(normal.font, normal.element, size=Pt(10))
    for section in doc.sections:
        section.left_margin = section.right_margin = Pt(42)

    # ------------------------------------------------------------
    # 문서 머리 — NCR과 같은 1행 2열 제목 표 + 작성정보 표
    # ------------------------------------------------------------
    title_t = doc.add_table(rows=1, cols=2)
    title_t.style = "Table Grid"
    tc0, tc1 = title_t.rows[0].cells
    _shade_cell(tc0, NAVY_HEX)
    p0 = tc0.paragraphs[0]
    r0 = p0.add_run("오늘의 현황 보고서")
    _apply_korean(r0.font, r0._element, size=Pt(14), bold=True)
    r0.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    r0b = p0.add_run("  Daily Status Report")
    _apply_korean(r0b.font, r0b._element, size=Pt(9))
    r0b.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    _cell(tc1, "AI 용접부 RT 결함 자동선별 시스템 — '오늘의 현황' 화면 기준", size=9)

    info = _table(doc, rows=2, cols=4, widths=[90, 174, 90, 174])
    _cell(info.cell(0, 0), "작성일시", bold=True)
    _cell(info.cell(0, 1), now.strftime("%Y-%m-%d %H:%M"))
    _cell(info.cell(0, 2), "보고번호", bold=True)
    _cell(info.cell(0, 3), report_no)
    th_cells = info.rows[1].cells
    th_val = th_cells[1].merge(th_cells[2]).merge(th_cells[3])
    _cell(th_cells[0], "적용 임계값", bold=True)
    _cell(
        th_val,
        f"무결함 확신 ≥ {thresholds['nd_confident']:.2f}   ·   주의 기준 ≥ {thresholds['attention_t']:.2f}   ·   "
        f"확정 기준 ≥ {thresholds['confident_t']:.2f}   ·   margin ≥ {thresholds['margin_threshold']:.0f}",
    )
    for r in range(2):
        for c in (0, 2):
            _shade_cell(info.cell(r, c), STRIPE_HEX)

    # ------------------------------------------------------------
    # 1. 오늘 처리현황
    # ------------------------------------------------------------
    _bold_line(doc, "1  오늘 처리현황")
    total_n = sum(r["n"] for r in routing_summary)
    t1 = _table(doc, rows=len(routing_summary) + 2, cols=3)
    _head_row(t1, ["구분", "건수", "비율"])
    for i, row in enumerate(routing_summary, start=1):
        _body_cell(t1, i, 0, row["label"], bold=True)
        _body_cell(t1, i, 1, f"{row['n']:,}건")
        _body_cell(t1, i, 2, f"{row['pct']:.1%}")
    last = len(routing_summary) + 1
    _body_cell(t1, last, 0, "총 검사 물량", bold=True)
    _body_cell(t1, last, 1, f"{total_n:,}건", bold=True)
    _body_cell(t1, last, 2, "100%", bold=True)

    # ------------------------------------------------------------
    # 2. 핵심 지표
    # ------------------------------------------------------------
    _bold_line(doc, "2  핵심 지표")
    d1_miss = core_kpis[1][1]
    d4_miss = core_kpis[2][1]
    t2 = _table(doc, rows=2, cols=4)
    _head_row(t2, ["자동화율", "사람확인 대기", "D1(균열) 미검출", "D4(용입불량) 미검출"])
    _cell(t2.cell(1, 0), f"{core_kpis[0][1]:.1%}", bold=True, size=12, align_center=True)
    _cell(t2.cell(1, 1), f"{len(open_items)}건", bold=True, size=12, align_center=True)
    _cell(t2.cell(1, 2), f"{d1_miss:.1%}", bold=True, size=12, align_center=True)
    _cell(t2.cell(1, 3), f"{d4_miss:.1%}", bold=True, size=12, align_center=True)
    t2.cell(1, 2).paragraphs[0].runs[0].font.color.rgb = OK_RGB if d1_miss == 0 else BAD_RGB
    t2.cell(1, 3).paragraphs[0].runs[0].font.color.rgb = OK_RGB if d4_miss == 0 else BAD_RGB

    # ------------------------------------------------------------
    # 3. 검사자 처리 결과
    # ------------------------------------------------------------
    _bold_line(doc, "3  검사자 처리 결과")
    t3 = _table(doc, rows=2, cols=3)
    _head_row(t3, ["승인(양품)", "반려(불량)", "처리 완료"])
    _cell(t3.cell(1, 0), f"{approved_n}건", bold=True, size=12, align_center=True)
    _cell(t3.cell(1, 1), f"{rejected_n}건", bold=True, size=12, align_center=True)
    _cell(t3.cell(1, 2), f"{approved_n + rejected_n}건", bold=True, size=12, align_center=True)

    _plain_line(doc, f"발행된 부적합보고서(NCR) — {len(ncr_list)}건", size=10)
    if ncr_list:
        t3b = _table(doc, rows=len(ncr_list) + 1, cols=3, widths=[40, 120, 368])
        _head_row(t3b, ["No", "이미지 ID", "NCR 파일명"])
        for i, (img_id, fname) in enumerate(ncr_list, start=1):
            _body_cell(t3b, i, 0, str(i))
            _body_cell(t3b, i, 1, img_id)
            _body_cell(t3b, i, 2, fname)
    else:
        _plain_line(doc, "오늘 발행된 NCR이 없습니다.", size=9)

    # ------------------------------------------------------------
    # 4. 위험도순 확인 대기 목록
    # ------------------------------------------------------------
    _bold_line(doc, f"4  위험도순 확인 대기 목록 — {len(open_items)}건")
    if open_items:
        t4 = _table(doc, rows=len(open_items) + 1, cols=5, widths=[40, 80, 90, 50, 268])
        _head_row(t4, ["순위", "이미지 ID", "의심유형", "확률", "판정상태"])
        for i, item in enumerate(open_items, start=1):
            _body_cell(t4, i, 0, str(i))
            _body_cell(t4, i, 1, str(item["image_id"]))
            _body_cell(t4, i, 2, str(item["dominant_class"]))
            _body_cell(t4, i, 3, f"{item['calibrated_prob']:.2f}")
            _body_cell(t4, i, 4, STATUS_LABELS.get(item["status"], str(item["status"])))
    else:
        _plain_line(doc, "대기 케이스 없음", size=10)

    # ------------------------------------------------------------
    # 5. SPC 관리도 상태
    # ------------------------------------------------------------
    _bold_line(doc, "5  SPC 관리도 상태")
    t5 = _table(doc, rows=2, cols=5)
    _head_row(t5, ["최근 일자", "최근 불량률", "중심선(CL)", "관리상한(UCL)", "관리한계 이탈"])
    _cell(t5.cell(1, 0), spc["date"].strftime("%Y-%m-%d"), align_center=True)
    _cell(t5.cell(1, 1), f"{spc['value']:.1%}", bold=True, align_center=True)
    _cell(t5.cell(1, 2), f"{spc['cl']:.1%}", align_center=True)
    _cell(t5.cell(1, 3), f"{spc['ucl']:.1%}", align_center=True)
    breach_text = "이탈" if spc["breached"] else "정상"
    if spc.get("demo"):
        breach_text += " (시연)"
    _cell(t5.cell(1, 4), breach_text, bold=True, align_center=True)
    t5.cell(1, 4).paragraphs[0].runs[0].font.color.rgb = BAD_RGB if spc["breached"] else OK_RGB
    _plain_line(doc, f"관리도 표시 기간 내 UCL 초과 이상점: {spc['n_outliers']}건", size=9)

    foot = _plain_line(
        doc,
        "본 보고서는 대시보드 '오늘의 현황' 화면에 표시된 값을 그대로 자동 생성한 것입니다. "
        "AI 용접부 RT 자동선별 시스템 — 포커스팀",
        size=8, italic=True,
    )
    foot.paragraph_format.space_before = Pt(10)
    foot.alignment = WD_ALIGN_PARAGRAPH.LEFT

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue(), f"today_status_report_{now.strftime('%Y-%m-%d')}.docx"
